"""Independent reference for interest, payback and APR.

Written separately from the page's JavaScript engine (different language,
different author pass) so the two can check each other. It produces:
  tests/fixtures/loan-cases.json   expected numbers for the automated tests
  reference/loan-reference.xlsx    a spreadsheet with live formulas
                                   (ROUND, PMT, IPMT, FV, XIRR) you can open in
                                   Excel, Numbers or Google Sheets.

Rules both sides follow (also shown on the page):
  - One period = one calendar month, starting on the start date.
  - Interest per period, by day-count rule:
      "30/360" (preset): balance x yearly rate / 12
      "act/365": balance x yearly rate x days in period / 365
      "act/360": balance x yearly rate x days in period / 360
  - Every period's interest is rounded to the cent.
  - interestOnly: interest paid monthly, full loan repaid in the last month.
  - bullet: nothing paid until the end; interest is added up monthly.
            compounding = monthly means interest is also charged on added interest.
  - amortizing: equal monthly payment = ROUND(PMT(rate/12, months left, -balance), 2),
                recalculated when the rate changes; the last payment clears the rest.
  - Origination fee reduces the money you actually receive.
  - Effective APR = XIRR of: money received at start (+), every payment (-).
Usage: .tools/py/bin/python reference/reference_loan.py
"""
import calendar
import datetime as dt
import json
import math
import os

from openpyxl import Workbook
from openpyxl.styles import Font

HERE = os.path.dirname(__file__)


def r2(x):
    """Round to cents, halves away from zero (like spreadsheet ROUND)."""
    s = 1 if x >= 0 else -1
    return s * math.floor(abs(x) * 100 + 0.5 + 1e-9) / 100


def add_months(d, n):
    m = d.month - 1 + n
    y = d.year + m // 12
    m = m % 12 + 1
    day = min(d.day, calendar.monthrange(y, m)[1])
    return dt.date(y, m, day)


def rate_for(case, i):
    path = case["rate"]["path"] if case["rate"]["type"] == "variable" else [{"fromMonth": 0, "annual": case["rate"]["annual"]}]
    r = path[0]["annual"]
    for step in path:
        if step["fromMonth"] <= i:
            r = step["annual"]
    return r


def period_interest(balance, annual, days, day_count):
    if day_count == "30/360":
        return balance * annual / 12
    if day_count == "act/365":
        return balance * annual * days / 365
    if day_count == "act/360":
        return balance * annual * days / 360
    raise ValueError(day_count)


def pmt(rate, n, pv):
    if rate == 0:
        return pv / n
    # 1 - (1 + r)^-n via expm1/log1p: stays accurate for rates close to 0
    return pv * rate / -math.expm1(-n * math.log1p(rate))


def schedule(case):
    start = dt.date.fromisoformat(case["startDate"])
    n = case["durationMonths"]
    bal = case["loanAmount"]
    accrued = 0.0
    rows = []
    payment = None
    last_rate = None
    for i in range(n):
        d0, d1 = add_months(start, i), add_months(start, i + 1)
        days = (d1 - d0).days
        annual = rate_for(case, i)
        rep = case["repayment"]
        if rep == "interestOnly":
            interest = r2(period_interest(bal, annual, days, case["dayCount"]))
            principal = bal if i == n - 1 else 0.0
            pay = r2(interest + principal)
            rows.append(dict(month=i + 1, date=d1.isoformat(), rate=annual, days=days, interest=interest, principal=r2(principal), payment=pay, balanceAfter=r2(bal - principal)))
            bal = r2(bal - principal)
        elif rep == "bullet":
            base = bal + (accrued if case["compounding"] == "monthly" else 0.0)
            interest = r2(period_interest(base, annual, days, case["dayCount"]))
            accrued = r2(accrued + interest)
            last = i == n - 1
            pay = r2(bal + accrued) if last else 0.0
            rows.append(dict(month=i + 1, date=d1.isoformat(), rate=annual, days=days, interest=interest, principal=r2(bal) if last else 0.0, payment=pay, balanceAfter=0.0 if last else r2(bal + accrued)))
        elif rep == "amortizing":
            if payment is None or annual != last_rate:
                payment = r2(pmt(annual / 12, n - i, bal))
                last_rate = annual
            interest = r2(period_interest(bal, annual, days, case["dayCount"]))
            if i == n - 1:
                pay = r2(bal + interest)
            else:
                pay = payment
            principal = r2(pay - interest)
            bal = r2(bal - principal)
            rows.append(dict(month=i + 1, date=d1.isoformat(), rate=annual, days=days, interest=interest, principal=principal, payment=pay, balanceAfter=bal))
    return rows


def fee_amount(case):
    f = case["originationFee"]
    return r2(case["loanAmount"] * f["value"]) if f["type"] == "percent" else r2(f["value"])


def xirr(flows):
    """flows: list of (date, amount). Bisection on the XIRR equation."""
    t0 = flows[0][0]

    def npv(r):
        return sum(a / (1 + r) ** ((d - t0).days / 365) for d, a in flows)

    lo, hi = -0.99, 10.0
    for _ in range(300):
        mid = (lo + hi) / 2
        if npv(lo) * npv(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


def expected(case):
    rows = schedule(case)
    fee = fee_amount(case)
    total_interest = r2(sum(r["interest"] for r in rows))
    total_payback = r2(sum(r["payment"] for r in rows))
    received = r2(case["loanAmount"] - fee)
    start = dt.date.fromisoformat(case["startDate"])
    flows = [(start, received)] + [(dt.date.fromisoformat(r["date"]), -r["payment"]) for r in rows if r["payment"]]
    return dict(
        schedule=rows,
        originationFee=fee,
        amountReceived=received,
        totalInterest=total_interest,
        totalPayback=total_payback,
        totalCost=r2(total_payback - case["loanAmount"] + fee),
        interestShareOfPayback=total_interest / total_payback,
        effectiveApr=xirr(flows),
    )


def base(**kw):
    c = dict(loanAmount=10000.0, durationMonths=12, rate=dict(type="fixed", annual=0.10), repayment="interestOnly",
             compounding="none", dayCount="30/360", startDate="2026-01-01",
             originationFee=dict(type="percent", value=0.0))
    c.update(kw)
    return c


CASES = {
    "interest-only, 30/360": base(),
    "interest-only, act/365": base(dayCount="act/365"),
    "interest-only, act/360, 1% fee": base(dayCount="act/360", originationFee=dict(type="percent", value=0.01)),
    "bullet, simple": base(repayment="bullet"),
    "bullet, monthly compounding": base(repayment="bullet", compounding="monthly"),
    "bullet, compounding, 24 months, act/365": base(repayment="bullet", compounding="monthly", durationMonths=24, dayCount="act/365"),
    "amortizing, 12 months": base(repayment="amortizing"),
    "amortizing, 36 months, 250 fixed fee": base(repayment="amortizing", durationMonths=36, loanAmount=25000.0, rate=dict(type="fixed", annual=0.0725), originationFee=dict(type="fixed", value=250.0)),
    "interest-only, variable 8% then 12%": base(rate=dict(type="variable", path=[dict(fromMonth=0, annual=0.08), dict(fromMonth=6, annual=0.12)])),
    "amortizing, variable 6% then 9%": base(repayment="amortizing", durationMonths=24, rate=dict(type="variable", path=[dict(fromMonth=0, annual=0.06), dict(fromMonth=12, annual=0.09)])),
    "zero interest, 2% fee": base(rate=dict(type="fixed", annual=0.0), originationFee=dict(type="percent", value=0.02)),
    "start on 31st (month-end handling)": base(startDate="2026-01-31", dayCount="act/365"),
}


def write_xlsx(results):
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    for idx, (name, case) in enumerate(CASES.items(), 1):
        ws = wb.create_sheet(f"Case {idx}")
        ws["A1"] = name; ws["A1"].font = bold
        inputs = [("Loan amount", case["loanAmount"]), ("Months", case["durationMonths"]),
                  ("Start date", dt.date.fromisoformat(case["startDate"])), ("Day count", case["dayCount"]),
                  ("Repayment", case["repayment"]), ("Compounding", case["compounding"]),
                  ("Origination fee", fee_amount(case))]
        for r, (k, v) in enumerate(inputs, 3):
            ws.cell(r, 1, k); ws.cell(r, 2, v)
        # B3 loan, B4 months, B5 start, B9 fee
        hdr = ["Month", "Period start", "Period end", "Days", "Yearly rate", "Balance at start", "Interest", "Principal", "Payment", "Balance after"]
        top = 12
        for c, h in enumerate(hdr, 1):
            ws.cell(top, c, h).font = bold
        n = case["durationMonths"]
        rows = results[name]["schedule"]
        for i in range(n):
            r = top + 1 + i
            ws.cell(r, 1, i + 1)
            ws.cell(r, 2, f"=EDATE($B$5,{i})")
            ws.cell(r, 3, f"=EDATE($B$5,{i + 1})")
            ws.cell(r, 4, f"=C{r}-B{r}")
            ws.cell(r, 5, rows[i]["rate"])
            ws.cell(r, 6, "=$B$3" if i == 0 else f"=J{r - 1}")
            dc = case["dayCount"]
            factor = "/12" if dc == "30/360" else (f"*D{r}/365" if dc == "act/365" else f"*D{r}/360")
            if case["repayment"] == "bullet" and case["compounding"] == "monthly":
                ws.cell(r, 7, f"=ROUND(F{r}*E{r}{factor},2)")
            else:
                ws.cell(r, 7, f"=ROUND($B$3*E{r}{factor},2)" if case["repayment"] == "bullet" else f"=ROUND(F{r}*E{r}{factor},2)")
            last = i == n - 1
            rep = case["repayment"]
            if rep == "interestOnly":
                ws.cell(r, 8, f"=F{r}" if last else 0)
                ws.cell(r, 9, f"=G{r}+H{r}")
                ws.cell(r, 10, f"=F{r}-H{r}")
            elif rep == "bullet":
                ws.cell(r, 8, "=$B$3" if last else 0)
                ws.cell(r, 9, f"=F{r}+G{r}" if last else 0)
                ws.cell(r, 10, 0 if last else f"=F{r}+G{r}")
            else:
                changed = i == 0 or rows[i]["rate"] != rows[i - 1]["rate"]
                pmt_cell = f"ROUND(PMT(E{r}/12,{n - i},-F{r}),2)" if changed else f"I{r - 1}"
                ws.cell(r, 9, f"=F{r}+G{r}" if last else f"={pmt_cell}")
                ws.cell(r, 8, f"=I{r}-G{r}")
                ws.cell(r, 10, f"=F{r}-H{r}")
        end = top + n
        s = end + 2
        ws.cell(s, 1, "Total interest").font = bold; ws.cell(s, 2, f"=SUM(G{top + 1}:G{end})")
        ws.cell(s + 1, 1, "Total payback").font = bold; ws.cell(s + 1, 2, f"=SUM(I{top + 1}:I{end})")
        ws.cell(s + 2, 1, "Interest share of payback").font = bold; ws.cell(s + 2, 2, f"=B{s}/B{s + 1}")
        ws.cell(s + 3, 1, "Money received (loan - fee)").font = bold; ws.cell(s + 3, 2, "=B3-B9")
        # XIRR cash flows in columns L:M
        ws.cell(top, 12, "Cash-flow date").font = bold; ws.cell(top, 13, "Cash flow").font = bold
        ws.cell(top + 1, 12, "=B5"); ws.cell(top + 1, 13, f"=B{s + 3}")
        for i in range(n):
            r = top + 2 + i
            ws.cell(r, 12, f"=C{top + 1 + i}"); ws.cell(r, 13, f"=-I{top + 1 + i}")
        ws.cell(s + 4, 1, "Effective APR (XIRR)").font = bold
        ws.cell(s + 4, 2, f"=XIRR(M{top + 1}:M{top + 1 + n},L{top + 1}:L{top + 1 + n})")
        if case["repayment"] == "bullet" and case["compounding"] == "monthly" and case["dayCount"] == "30/360":
            ws.cell(s + 5, 1, "Check: FV formula (no rounding)"); ws.cell(s + 5, 2, f"=FV(E{top + 1}/12,B4,0,-B3)-B3")
        if case["repayment"] == "amortizing" and case["rate"]["type"] == "fixed" and case["dayCount"] == "30/360":
            ws.cell(s + 5, 1, "Check: IPMT month 1"); ws.cell(s + 5, 2, f"=IPMT(E{top + 1}/12,1,B4,-B3)")
        ws.cell(s + 7, 1, "Expected by the page (Python reference):").font = bold
        e = results[name]
        for k, label in enumerate(["totalInterest", "totalPayback", "effectiveApr"]):
            ws.cell(s + 8 + k, 1, label); ws.cell(s + 8 + k, 2, e[label])
        for col, w in zip("ABCDEFGHIJKLM", [34, 13, 13, 6, 11, 15, 10, 11, 11, 14, 2, 15, 12]):
            ws.column_dimensions[col].width = w
        for r in range(top + 1, end + 1):
            for c in (2, 3):
                ws.cell(r, c).number_format = "yyyy-mm-dd"
        ws["B5"].number_format = "yyyy-mm-dd"
        for r in range(top + 2, top + 2 + n):
            ws.cell(r, 12).number_format = "yyyy-mm-dd"
        ws.cell(top + 1, 12).number_format = "yyyy-mm-dd"
    wb.save(os.path.join(HERE, "loan-reference.xlsx"))


if __name__ == "__main__":
    results = {name: expected(c) for name, c in CASES.items()}
    os.makedirs(os.path.join(HERE, "..", "tests", "fixtures"), exist_ok=True)
    with open(os.path.join(HERE, "..", "tests", "fixtures", "loan-cases.json"), "w") as f:
        json.dump([dict(name=n, terms=CASES[n], expected=results[n]) for n in CASES], f, indent=1)
    write_xlsx(results)
    for n, e in results.items():
        print(f"{n:42s} interest {e['totalInterest']:>10.2f}  payback {e['totalPayback']:>10.2f}  APR {e['effectiveApr']*100:7.4f}%")
